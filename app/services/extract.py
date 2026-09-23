"""Turn an uploaded file into either native text or page images.

Most CVs arriving by email are digital PDFs with a real text layer. Running OCR
on those is slower AND less accurate, so we check first and skip OCR entirely.
"""
import io
import logging

import fitz  # PyMuPDF
import numpy as np

from app.core.config import settings

log = logging.getLogger(__name__)

PDF_MIMES = {"application/pdf"}
IMAGE_MIMES = {"image/png", "image/jpeg", "image/jpg", "image/tiff", "image/webp"}
DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def pdf_text_layer(data: bytes, min_chars: int) -> list[str] | None:
    """Return per-page text if the PDF has a usable text layer, else None."""
    with fitz.open(stream=data, filetype="pdf") as doc:
        pages = [p.get_text("text") for p in doc]
    total = sum(len(t.strip()) for t in pages)
    return pages if total >= min_chars else None


def pdf_to_images(data: bytes, dpi: int, max_pages: int) -> list[np.ndarray]:
    images = []
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc[:max_pages]:
            pix = page.get_pixmap(matrix=matrix, colorspace=fitz.csRGB)
            arr = np.frombuffer(pix.samples, dtype=np.uint8)
            images.append(arr.reshape(pix.height, pix.width, 3))
    return images


def image_to_array(data: bytes) -> np.ndarray:
    from PIL import Image

    img = Image.open(io.BytesIO(data)).convert("RGB")
    return np.array(img)


def docx_text(data: bytes) -> list[str]:
    """DOCX CVs are common enough to be worth handling natively."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return ["\n".join(parts)]


def prepare(data: bytes, mime: str, filename: str) -> dict:
    """-> {"kind": "text"|"images", "pages": [...], "page_count": n}"""
    mime = (mime or "").lower()
    lower = filename.lower()

    if mime in DOCX_MIMES or lower.endswith(".docx"):
        return {"kind": "text", "pages": docx_text(data), "page_count": 1}

    if mime in PDF_MIMES or lower.endswith(".pdf"):
        native = pdf_text_layer(data, settings.min_text_chars)
        if native is not None:
            return {
                "kind": "text",
                "pages": native[: settings.max_pages],
                "page_count": len(native),
            }
        images = pdf_to_images(data, settings.ocr_dpi, settings.max_pages)
        return {"kind": "images", "pages": images, "page_count": len(images)}

    if mime in IMAGE_MIMES or lower.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp")):
        return {"kind": "images", "pages": [image_to_array(data)], "page_count": 1}

    raise ValueError(f"Unsupported file type: {mime or filename}")
